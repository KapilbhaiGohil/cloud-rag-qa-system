import io
import fitz 
from docx import Document
from pptx import Presentation
from PIL import Image
import pytesseract

def parse_pdf(file_bytes: bytes) -> str:
    text = ""
    try:
        with fitz.open(stream=file_bytes, filetype="pdf") as doc:
            for page in doc:
                text += page.get_text() + "\n"
    except Exception as e:
        print(f"PDF parsing error: {e}")
    return text.strip()

def parse_docx(file_bytes: bytes) -> str:
    text = ""
    try:
        doc = Document(io.BytesIO(file_bytes))
        text = "\n".join([para.text for para in doc.paragraphs])
    except Exception as e:
        print(f"DOCX parsing error: {e}")
    return text.strip()

def parse_pptx(file_bytes: bytes) -> str:
    text = ""
    try:
        prs = Presentation(io.BytesIO(file_bytes))
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text += shape.text + "\n"
    except Exception as e:
        print(f"PPTX parsing error: {e}")
    return text.strip()

def parse_image(file_bytes: bytes) -> str:
    text = ""
    try:
        image = Image.open(io.BytesIO(file_bytes))
        text = pytesseract.image_to_string(image)
    except Exception as e:
        print(f"Image OCR error: {e}")
    return text.strip()

def parse_txt(file_bytes: bytes) -> str:
    try:
        return file_bytes.decode('utf-8', errors='ignore')
    except Exception as e:
        print(f"TXT parsing error: {e}")
        return ""

def extract_text_from_file(file_bytes: bytes, filename: str, content_type: str) -> str:
    filename_lower = filename.lower()
    
    if content_type == "application/pdf" or filename_lower.endswith(".pdf"):
        return parse_pdf(file_bytes)
        
    elif "wordprocessingml.document" in content_type or filename_lower.endswith(".docx"):
        return parse_docx(file_bytes)
        
    elif "presentationml.presentation" in content_type or filename_lower.endswith(".pptx"):
        return parse_pptx(file_bytes)
        
    elif content_type.startswith("image/") or filename_lower.endswith((".png", ".jpg", ".jpeg")):
        return parse_image(file_bytes)
        
    else:
        return parse_txt(file_bytes)